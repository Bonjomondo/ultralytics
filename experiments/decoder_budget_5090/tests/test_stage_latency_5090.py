"""Stage attribution must remain tied to complete, equivalent measurements."""
import copy
import csv
import json
from pathlib import Path

import pytest

from experiments.decoder_budget_5090 import report_5090 as report
from experiments.decoder_budget_5090 import stage_latency_5090 as stage


@pytest.fixture
def raw():
    return json.loads((stage.ROOT / stage.EVIDENCE).read_text())


def records(archive):
    return [{"variant": v, "role": "paired_150e_screen", "training_status": "completed", "run_id": "paired/" + v,
             "checkpoint_sha256": {"best": m["checkpoint_sha256"]}} for v, m in archive["models"].items()]


def test_real_evidence_recomputes_stage_shares_and_round_ranges(raw):
    result = stage.summarize(raw)
    for model in result.values():
        for item in model["summary"].values():
            assert item["bare_whole"]["samples"] == 250
            assert sum(s["fraction_of_staged_forward"] for s in item["stages"].values()) == pytest.approx(1, abs=.0001)
            for entry in item["stages"].values():
                assert entry["round_mean_min_ms"] <= entry["mean_ms"] <= entry["round_mean_max_ms"]
    assert result["c2"]["summary"]["300"]["stages"]["backbone_0_to_4"]["mean_ms"] == pytest.approx(2.9213811)


@pytest.mark.parametrize("corruption", ["duplicate_round", "truncated_samples", "nan", "stage_accounting", "false_summary", "identity", "hook"])
def test_corrupt_stage_evidence_is_rejected(raw, corruption):
    model = raw["models"]["c0"]
    block = next(b for b in model["blocks"] if b["mode"] == "staged")
    if corruption == "duplicate_round":
        model["blocks"][0] = copy.deepcopy(model["blocks"][1])
    elif corruption == "truncated_samples":
        block["measurements"]["whole"].pop()
    elif corruption == "nan":
        block["measurements"]["whole"][0] = float("nan")
    elif corruption == "stage_accounting":
        block["measurements"][stage.STAGES[0]][0] += 1
    elif corruption == "false_summary":
        model["summary"]["300"]["bare_whole"]["mean_ms"] += 1
    elif corruption == "identity":
        model["identity_checks"]["150"] = False
    else:
        model["hooks_remaining"] = 1
    with pytest.raises(ValueError):
        stage.summarize(raw)


def test_unrelated_weights_are_not_assigned_this_assessment():
    archive = stage.load_archive()
    training = records(archive)
    attached = stage.collect(training)
    assert attached["models"]["c2"]["run_id"] == "paired/c2"
    for record in training:
        record["checkpoint_sha256"]["best"] = "other-weights"
    assert stage.collect(training) is None


@pytest.mark.parametrize("failure", ["one_checkpoint_changed", "failed_training", "duplicate_training"])
def test_partially_matching_or_ambiguous_training_cannot_claim_paired_assessment(failure):
    training = records(stage.load_archive())
    if failure == "one_checkpoint_changed":
        training[0]["checkpoint_sha256"]["best"] = "other-weights"
    elif failure == "failed_training":
        training[0]["training_status"] = "failed"
    else:
        training.append(copy.deepcopy(training[0]))
    with pytest.raises(ValueError):
        stage.collect(training)


def test_archive_digest_prevents_raw_replacement(tmp_path):
    path = tmp_path / stage.EVIDENCE
    path.parent.mkdir(parents=True)
    path.write_text("{}")
    identity = tmp_path / stage.IDENTITY
    identity.write_bytes((stage.ROOT / stage.IDENTITY).read_bytes())
    with pytest.raises(ValueError, match="SHA mismatch"):
        stage.load_archive(tmp_path)


def test_stage_outputs_replay_without_server_or_archive_access(tmp_path, monkeypatch):
    archive = stage.collect(records(stage.load_archive()))
    data = {"schema": "5090.experiments-report.v1", "generated_at": "archived-time", "records": [],
            "collection_errors": [], "stage_assessment": archive}
    def no_archive(*args, **kwargs):
        raise AssertionError("Offline rendering must not access the archive")
    monkeypatch.setattr(stage, "load_archive", no_archive)
    output = tmp_path / "offline"
    report.generate_reports(tmp_path / "missing-runs", output, data=data)
    saved = json.loads((output / "experiments.json").read_text())
    assert saved["stage_assessment"]["assessment_sha256"] == archive["assessment_sha256"]
    text = (output / "experiments.md").read_text()
    assert "39.3%" in text and "访存或CPU提交" in text and "不替换上方200次整网计时" in text
    with (output / "stage_latency.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 30 and float(rows[0]["fraction_of_staged_forward"]) < 1
    assert {r["queries"] for r in rows} == {"300", "200", "150"}
    from openpyxl import load_workbook
    sheet = load_workbook(output / "experiments.xlsx")["Stage latency"]
    assert sheet.max_row == 31
    column = stage.FIELDS.index("fraction_of_staged_forward") + 1
    assert sheet.cell(2, column).number_format == "0.00%"
