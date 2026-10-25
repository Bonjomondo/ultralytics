"""Reports must distinguish full validation, engineering gates and stale metrics."""
import argparse
import csv
import json
from pathlib import Path

import pytest

from experiments.decoder_budget_5090 import queue_5090
from experiments.decoder_budget_5090 import report_5090 as report


def write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data))


def evidence(root, *, name="5090-c0", status="completed", scope="full_val", role="paired_150e_screen"):
    run = root / name
    write(run / "experiment_5090.json", {"experiment_name": name, "variant": "c0", "status": status,
          "experiment_role": role, "requested_epochs": 150, "current_epoch": 150,
          "checkpoint_sha256": {"best": "weights-a"}})
    (run / "results.csv").write_text("epoch,metrics/mAP50-95(B)\n" + "".join(f"{i},0.2\n" for i in range(1, 151)))
    sweep = run.with_name(run.name + "-query-full")
    protocol = {"scope": scope, "checkpoint_sha256": "weights-a", "annotation_sha256": "gt-a"}
    conditions = []
    for budget, ap in ((300, .2), (250, .198)):
        query = sweep / f"q{budget}"
        write(query / "manifest.json", {"status": "completed", "image_count": 548,
              "provenance": {**protocol, "query_budget": budget},
              "identity_verification": "exact_native_observed_restored",
              "batches": [{"identity_passed": True, "restoration_passed": True}]})
        write(query / "coverage_summary.json", {"rows": [{"budget": budget, "group": "all", "iou_threshold": .5,
              "coverage": .8, "class_correct_coverage": .7}]})
        conditions.append({"queries": budget, "images": 548, "metrics": {"AP": ap, "APs": ap / 2, "APl": -1},
                           "manifest_sha256": report.sha(query / "manifest.json"),
                           "coverage_summary_sha256": report.sha(query / "coverage_summary.json"),
                           "input_gt_identity_sha256": "same-images-and-gt"})
    write(sweep / "summary.json", {"status": "completed", "protocol": protocol, "conditions": conditions})
    return run, sweep


def test_full_results_compute_only_same_checkpoint_delta_and_keep_raw_units(tmp_path):
    evidence(tmp_path / "5090-a")
    data = report.collect(tmp_path)
    main = report.rows(data, main=True)
    assert len(main) == 2 and main[0]["AP"] == .2
    assert main[1]["delta_AP_pp"] == pytest.approx(-.2)
    assert main[0]["APl"] is None  # unavailable COCO area is not a negative AP
    assert "20.00" in report.markdown(data)


@pytest.mark.parametrize("status", ["failed", "interrupted", "running"])
def test_stale_evaluations_never_enter_main_results(tmp_path, status):
    evidence(tmp_path / "5090-a", status=status)
    data = report.collect(tmp_path)
    assert report.rows(data, main=True) == []
    assert all(row["evaluation_status"] == "stale" and row["AP"] is None for row in report.rows(data))


@pytest.mark.parametrize("scope,role", [("subset_engineering_check", "paired_150e_screen"),
                                       ("full_val", "engineering_smoke_not_architecture_result")])
def test_subset_and_short_runs_are_kept_outside_main_results(tmp_path, scope, role):
    evidence(tmp_path / "5090-a", scope=scope, role=role)
    data = report.collect(tmp_path)
    assert not report.rows(data, main=True) and len(report.rows(data)) == 2


@pytest.mark.parametrize("failure", ["checkpoint", "manifest", "coverage", "input_identity", "partial", "epochs"])
def test_incompatible_or_incomplete_evidence_is_not_reported_as_valid_ap(tmp_path, failure):
    run, sweep = evidence(tmp_path / "5090-a")
    path = sweep / "summary.json"
    summary = report.read(path)
    if failure == "checkpoint":
        summary["protocol"]["checkpoint_sha256"] = "different-weights"
    elif failure == "manifest":
        (sweep / "q300/manifest.json").write_text("{}")
    elif failure == "coverage":
        (sweep / "q300/coverage_summary.json").write_text("{}")
    elif failure == "input_identity":
        summary["conditions"][1]["input_gt_identity_sha256"] = "different-inputs"
    elif failure == "partial":
        summary["status"] = "running"
    else:
        (run / "results.csv").write_text("epoch,metrics/mAP50-95(B)\n1,0.2\n")
    write(path, summary)
    data = report.collect(tmp_path)
    valid = report.rows(data, main=True)
    if failure in ("manifest", "coverage"):
        assert len(valid) == 1 and valid[0]["delta_AP_pp"] is None
    else:
        assert not valid


def test_identical_run_names_in_different_cohorts_keep_separate_baselines(tmp_path):
    evidence(tmp_path / "5090-a")
    evidence(tmp_path / "5090-b")
    records = report.collect(tmp_path)
    result = report.rows(records, main=True)
    assert len({row["run_id"] for row in result}) == 2
    assert [row["delta_AP_pp"] for row in result].count(0) == 2


def test_saved_snapshot_regenerates_all_outputs_without_server_assets(tmp_path):
    root, output = tmp_path / "runs", tmp_path / "output"
    evidence(root / "5090-a")
    write(root / "5090-a/5090-c0-query-full/q300/benchmark.json",
          {"latency_ms": [1, 3, 8], "repetitions": 3, "median_ms": 3, "p95_ms": 8})
    data = report.collect(root)
    report.generate_reports(root, output, data=data)
    snapshot = report.read(output / "experiments.json")
    second = tmp_path / "offline"
    report.generate_reports(tmp_path / "missing-server-runs", second, data=snapshot)
    assert (second / "experiments.md").read_bytes() == (output / "experiments.md").read_bytes()
    with (output / "experiments.csv").open() as stream:
        row = next(csv.DictReader(stream))
        assert float(row["AP"]) == .2
        assert float(row["mean_ms"]) == 4 and float(row["fps_from_mean"]) == 250
    from openpyxl import load_workbook
    book = load_workbook(second / "experiments.xlsx")
    sheet = book["Query results"]
    index = list(report.FIELDS).index("AP") + 1
    assert sheet.cell(2, index).value == .2 and sheet.cell(2, index).number_format == "0.00%"
    assert sheet.cell(2, list(report.FIELDS).index("fps_from_mean") + 1).value == 250
    assert len(list((second / "plots").glob("*.png"))) == 3


def test_fps_uses_mean_latency_instead_of_median_or_average_reciprocals():
    measured = report.timing({"latency_ms": [1, 3, 8], "repetitions": 3, "median_ms": 3})
    assert measured["mean_ms"] == 4 and measured["fps_from_mean"] == 250
    assert measured["median_ms"] == 3
    assert measured["peak_allocated_mib"] is None


@pytest.mark.parametrize("samples", [[], [0], [-1], [float("nan")], [float("inf")]])
def test_invalid_latency_samples_never_generate_fps(samples):
    with pytest.raises(ValueError):
        report.timing({"latency_ms": samples})


def test_missing_or_corrupt_timing_preserves_ap_and_leaves_fps_unavailable(tmp_path):
    _, sweep = evidence(tmp_path / "5090-a")
    write(sweep / "q300/benchmark.json", {"latency_ms": [1, 3], "repetitions": 200})
    data = report.collect(tmp_path)
    result = report.rows(data, main=True)
    assert len(result) == 2 and result[0]["AP"] == .2
    assert result[0]["benchmark_status"] == "invalid"
    assert result[1]["benchmark_status"] == "not_measured"
    assert all(row["fps_from_mean"] is None for row in result)


def test_queue_refreshes_after_state_changes_and_retains_report_error(tmp_path, monkeypatch):
    instance = queue_5090.Queue(argparse.Namespace(root=tmp_path / "5090-a", sdk=tmp_path / "sdk",
                               data_root=tmp_path / "data", report_dir=tmp_path / "reports"))
    instance.root.mkdir()
    instance.state = {"status": "running"}
    observed = []
    def generate(runs, output):
        observed.append((runs, output))
        return {"generated_at": "now", "collection_errors": []}
    monkeypatch.setattr(report, "generate_reports", generate)
    instance.update(phase="5090-training")
    assert observed and report.read(instance.status)["reporting"]["status"] == "completed"
    def fail(*args):
        raise RuntimeError("report failure")
    monkeypatch.setattr(report, "generate_reports", fail)
    instance.update(phase="5090-next-phase")
    state = report.read(instance.status)
    assert state["phase"] == "5090-next-phase" and state["reporting"]["status"] == "failed"
    assert "report failure" in state["reporting"]["error"]
