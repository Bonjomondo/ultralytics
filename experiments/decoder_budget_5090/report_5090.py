"""Generate experiment reports from saved 5090 training and query-sweep evidence.

No model, checkpoint or detection tensor is loaded. JSON/CSV retain 0--1 metric
values; Markdown/Excel display percentages. Rendering from a saved JSON works
without the original server's runs, datasets, SDK or CUDA.
"""
import argparse
import csv
from datetime import datetime
import fcntl
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile
from zoneinfo import ZoneInfo

if __package__:
    from . import stage_latency_5090 as stage_latency
else:
    import stage_latency_5090 as stage_latency

ROOT = Path(__file__).resolve().parents[2]
METRICS = ("AP", "AP50", "AP75", "APs", "APm", "APl", "AR1", "AR10", "AR100", "ARs", "ARm", "ARl")
BEST_KEY = "metrics/mAP50-95(B)"
FIELDS = ("experiment", "run_id", "variant", "queries", "training_status", "evaluation_status", "epochs", "images",
          *METRICS, "delta_AP_pp", "delta_APs_pp", "mean_ms", "median_ms", "p95_ms", "fps_from_mean", "peak_allocated_mib",
          "benchmark_status", "benchmark_reason",
          "coverage_05", "class_correct_coverage_05", "identity_passed", "scope", "reason")


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return None
    return float(value)


def valid_metric(value):
    value = number(value)
    return value if value is not None and 0 <= value <= 1 else None


def timing(benchmark):
    """Batch-one forward FPS is reciprocal mean latency, not mean(1/latency)."""
    if not isinstance(benchmark, dict):
        raise ValueError("计时记录必须为JSON对象")
    samples = benchmark.get("latency_ms", [])
    if not isinstance(samples, list) or not samples:
        raise ValueError("缺少原始计时样本，不能计算平均延迟/FPS")
    values = [number(value) for value in samples]
    if any(value is None or value <= 0 for value in values):
        raise ValueError("计时样本必须是有限正数")
    if benchmark.get("repetitions", len(values)) != len(values):
        raise ValueError("计时样本数与repetitions不一致")
    mean = math.fsum(values) / len(values)
    allocated = number(benchmark.get("peak_allocated_bytes"))
    return {"mean_ms": mean, "fps_from_mean": 1000 / mean,
            "median_ms": number(benchmark.get("median_ms")), "p95_ms": number(benchmark.get("p95_ms")),
            "peak_allocated_mib": allocated / 2 ** 20 if allocated is not None else None,
            "benchmark_scope": benchmark.get("scope"), "benchmark_status": "completed", "benchmark_reason": ""}


def curves(path):
    if not path.is_file():
        return []
    with path.open() as stream:
        return [{key.strip(): number(float(value)) for key, value in row.items()} for row in csv.DictReader(stream)]


def collect_training(path):
    state = read(path)
    run = path.parent
    curve = curves(run / "results.csv")
    if state.get("status") == "completed" and (len(curve) != state.get("requested_epochs")
            or state.get("current_epoch") != len(curve)):
        state.update(status="inconsistent", error="完成状态与训练曲线/请求轮数不一致")
    scored = [row for row in curve if row.get(BEST_KEY) is not None]
    best = max(scored, key=lambda row: row[BEST_KEY]) if scored else {}
    return {"experiment": state.get("experiment_name", run.name), "run_id": run.parent.name + "/" + run.name,
            "variant": state.get("variant", "unknown"),
            "run": str(run), "cohort": str(run.parent), "training_status": state.get("status", "unknown"),
            "role": state.get("experiment_role", "unknown"), "epochs": len(curve),
            "requested_epochs": state.get("requested_epochs"), "seed": state.get("seed"),
            "best_epoch": best.get("epoch"), "best_internal_map": best.get(BEST_KEY),
            "internal_precision": best.get("metrics/precision(B)"), "internal_recall": best.get("metrics/recall(B)"),
            "curve": curve, "environment": state.get("environment", {}),
            "checkpoints": state.get("checkpoints", {}), "checkpoint_sha256": state.get("checkpoint_sha256", {}),
            "source_commit": state.get("source_commit"), "data_yaml_sha256": state.get("data_yaml_sha256"),
            "shared_initialization_sha256": state.get("expected_shared_state_sha256"),
            "error": state.get("error"), "state_sha256": sha(path), "sweeps": []}


def blank_row(training, queries, scope, status, reason=""):
    row = {key: None for key in FIELDS}
    row.update(experiment=training["experiment"], run_id=training["run_id"], variant=training["variant"], queries=queries,
               training_status=training["training_status"], evaluation_status=status,
               epochs=training["epochs"], scope=scope, reason=reason,
               benchmark_status="not_measured", benchmark_reason="未保存计时结果")
    return row


def collect_sweep(path, training):
    summary_path = path / "summary.json"
    result = read(summary_path) if summary_path.is_file() else read(path / "progress.json") if (path / "progress.json").is_file() else {}
    protocol = result.get("protocol") or (read(path / "protocol.json") if (path / "protocol.json").is_file() else {})
    scope = protocol.get("scope", "unknown")
    sweep = {"run": str(path), "scope": scope, "status": result.get("status", "pending"),
             "protocol": protocol, "rows": [], "summary_sha256": sha(summary_path) if summary_path.is_file() else None}
    failure = read(path / "failure.json") if (path / "failure.json").is_file() else {}
    for condition in result.get("conditions", []):
        query = condition["queries"]
        row = blank_row(training, query, scope, "pending")
        row["images"] = condition.get("images")
        row["input_gt_identity_sha256"] = condition.get("input_gt_identity_sha256")
        row["observed_metrics"] = condition.get("metrics", {})
        try:
            if training["training_status"] != "completed":
                row.update(evaluation_status="stale", reason="训练未完成或已失败/中断，现有评估不进入主结果")
                sweep["rows"].append(row)
                continue
            if failure or result.get("status") != "completed":
                row.update(evaluation_status="failed" if failure else "partial", reason=failure.get("error", "四档评估尚未全部完成"))
                sweep["rows"].append(row)
                continue
            expected = training["checkpoint_sha256"].get("best")
            if not expected or protocol.get("checkpoint_sha256") != expected:
                raise ValueError("评估checkpoint SHA与训练best不一致或缺失")
            query_dir = path / f"q{query}"
            manifest_path = query_dir / "manifest.json"
            if sha(manifest_path) != condition["manifest_sha256"]:
                raise ValueError("观测manifest SHA不匹配")
            manifest = read(manifest_path)
            if manifest.get("status") != "completed" or manifest.get("image_count") != condition["images"]:
                raise ValueError("观测状态或图像数不匹配")
            provenance = manifest.get("provenance", {})
            if (provenance.get("checkpoint_sha256") != expected
                    or provenance.get("annotation_sha256") != protocol.get("annotation_sha256")
                    or provenance.get("query_budget") != query):
                raise ValueError("观测记录与评估的权重/标注/Query预算不一致")
            identity = manifest.get("identity_verification", "disabled") != "disabled" and bool(manifest["batches"]) and all(
                batch.get("identity_passed") and batch.get("restoration_passed") for batch in manifest["batches"])
            if not identity:
                raise ValueError("原生/观测/恢复等价门禁未通过")
            coverage_path = query_dir / "coverage_summary.json"
            if sha(coverage_path) != condition["coverage_summary_sha256"]:
                raise ValueError("覆盖分析SHA不匹配")
            coverage = next(item for item in read(coverage_path)["rows"]
                            if item["budget"] == query and item["group"] == "all" and item["iou_threshold"] == .5)
            row.update({metric: valid_metric(condition.get("metrics", {}).get(metric)) for metric in METRICS})
            if row["AP"] is None:
                raise ValueError("缺少有效正式AP")
            row.update(evaluation_status="completed", identity_passed=True,
                       coverage_05=valid_metric(coverage["coverage"]),
                       class_correct_coverage_05=valid_metric(coverage["class_correct_coverage"]))
            if (query_dir / "benchmark.json").is_file():
                try:
                    row.update(timing(read(query_dir / "benchmark.json")))
                except (OSError, ValueError, TypeError) as error:
                    row.update(benchmark_status="invalid", benchmark_reason=str(error))
        except (OSError, ValueError, KeyError, StopIteration, TypeError) as error:
            row.update({key: None for key in (*METRICS, "median_ms", "p95_ms", "peak_allocated_mib", "coverage_05", "class_correct_coverage_05")})
            row.update(evaluation_status="invalid", reason=str(error), identity_passed=False)
        sweep["rows"].append(row)
    complete = [row for row in sweep["rows"] if row["evaluation_status"] == "completed"]
    if complete and (len({row["input_gt_identity_sha256"] for row in complete}) != 1 or not complete[0]["input_gt_identity_sha256"]):
        for row in complete:
            row.update(evaluation_status="invalid", reason="跨Query输入/GT身份不一致")
            row.update({metric: None for metric in METRICS})
        complete = []
    baselines = [row for row in complete if row["queries"] == 300]
    if len(baselines) == 1:
        for row in complete:
            for metric in ("AP", "APs"):
                a, b = row[metric], baselines[0][metric]
                row[f"delta_{metric}_pp"] = 100 * (a - b) if a is not None and b is not None else None
    return sweep


def collect(runs_dir):
    runs_dir = Path(runs_dir).resolve()
    if not runs_dir.is_dir():
        raise FileNotFoundError(runs_dir)
    records, errors = [], []
    for path in sorted(runs_dir.rglob("experiment_5090.json")):
        try:
            record = collect_training(path)
            for sweep in sorted(path.parent.parent.glob(path.parent.name + "-query-*")):
                if sweep.is_dir():
                    record["sweeps"].append(collect_sweep(sweep, record))
            records.append(record)
        except (OSError, ValueError, KeyError, TypeError) as error:
            errors.append({"path": str(path), "error": str(error)})
    assessment = None
    try:
        assessment = stage_latency.collect(records)
    except (OSError, ValueError, KeyError, TypeError) as error:
        errors.append({"path": stage_latency.EVIDENCE, "error": str(error)})
    return {"schema": "5090.experiments-report.v1", "generated_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
            "runs_dir": str(runs_dir), "records": records, "collection_errors": errors,
            "stage_assessment": assessment,
            "units": "JSON/CSV metrics are 0–1; delta is percentage points; latency is milliseconds; fps_from_mean=1000/mean_ms for batch1 GPU forward"}


def rows(data, *, main=False):
    result = []
    for training in data["records"]:
        for sweep in training["sweeps"]:
            for row in sweep["rows"]:
                if not main or (training["role"] == "paired_150e_screen" and sweep["scope"] == "full_val"
                                and row["evaluation_status"] == "completed"):
                    result.append(row)
    return sorted(result, key=lambda row: (row["run_id"], -(row["queries"] or 0)))


def display(value, percent=False, digits=2):
    return "—" if value is None else f"{value * (100 if percent else 1):.{digits}f}"


def safe(value):
    return str(value).replace("|", "\\|").replace("\n", " ").replace("<", "&lt;")


def markdown(data):
    lines = ["# 5090 RT-DETR 实验汇总", "", f"更新时间：{data['generated_at']}", "",
             "[CSV](experiments.csv) · [JSON](experiments.json) · [Excel](experiments.xlsx) · [产物校验](MANIFEST.sha256)", "",
             "由 `experiments/decoder_budget_5090/report_5090.py` 自动扫描生成。正式COCO指标以百分数显示，Δ为相对同一checkpoint的300Q基线的百分点；JSON/CSV保留0–1原始指标。",
             "内部Precision/Recall/mAP仅描述训练日志。完整验证、8图门禁和1轮短跑分别展示；失败/中断后残留的评估标为stale，不进入主结果。",
             "不同checkpoint不直接计算Query压缩差值。单seed与单次计时不证明稳定收益；完整训练150轮不等于证明收敛。", "", "## 主结果：完整验证集", "",
             "| 实验 / Run | Q | 图像 | 状态 | AP | AP50 | AP75 | APs | APm | APl | ΔAP (pp) | ΔAPs (pp) | 平均GPU延迟 ms | GPU FPS |",
             "|---|---:|---:|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|"]
    main = rows(data, main=True)
    for row in main:
        values = [safe(row["run_id"]), str(row["queries"]), str(row["images"]), row["evaluation_status"],
                  *[display(row[key], True) for key in METRICS[:6]], display(row["delta_AP_pp"]), display(row["delta_APs_pp"]),
                  display(row.get("mean_ms"), digits=3), display(row.get("fps_from_mean"))]
        lines.append("| " + " | ".join(values) + " |")
    if not main:
        lines += ["", "尚无通过身份门禁的完整验证结果。"]
    lines += ["", "## 计时与候选覆盖", "", "计时为未融合FP32、batch1、50次预热/200次CUDA Event、无observer的GPU整网forward，不含传输/预处理/后处理；不是Decoder单独或端到端延迟。GPU FPS=1000/原始计时样本的平均延迟(ms)，不是中位延迟倒数或逐样本FPS的平均；没有计时样本时不估算FPS。Coverage为Encoder候选多对多GT覆盖，不是COCO召回/AP。", "",
              "| 实验 / Run | Q | 平均延迟 ms | 中位延迟 ms | P95 ms | GPU FPS | 峰值allocated MiB | Coverage@.5 | 正确类Coverage@.5 |",
              "|---|---:|---:|---:|---:|---:|---:|---:|---:|"]
    for row in main:
        lines.append("| " + " | ".join([safe(row["run_id"]), str(row["queries"]), display(row.get("mean_ms"), digits=3),
                                      display(row["median_ms"], digits=3), display(row["p95_ms"], digits=3),
                                      display(row.get("fps_from_mean")), display(row["peak_allocated_mib"]),
                                      display(row["coverage_05"], True), display(row["class_correct_coverage_05"], True)]) + " |")
    lines += stage_latency.markdown(data)
    lines += ["", "## 训练、门禁与异常状态", "", "| 实验 | 训练状态 | 已记录轮数 / 请求 | 最佳内部mAP轮次 | 最佳内部mAP (%) | 正式全量结果行数 |", "|---|---|---:|---:|---:|---:|"]
    for record in data["records"]:
        count = sum(row["run_id"] == record["run_id"] for row in main)
        lines.append(f"| {safe(record['run_id'])} | {safe(record['training_status'])} | {record['epochs']} / {record['requested_epochs']} | {display(record['best_epoch'], digits=0)} | {display(record['best_internal_map'], True)} | {count} |")
    lines += ["", "## 实验说明", ""]
    for record in data["records"]:
        lines += [f"### {safe(record['run_id'])}", "", f"角色：`{record['role']}`；训练状态：`{record['training_status']}`；seed：{record['seed']}。",
                  f"最佳日志轮次的内部Precision/Recall：{display(record['internal_precision'], True)}% / {display(record['internal_recall'], True)}%。",
                  f"GPU：{safe(record['environment'].get('gpu', 'unknown'))}；PyTorch：{safe(record['environment'].get('torch', 'unknown'))}；源码commit：`{record['source_commit']}`。",
                  "", f"服务器run：`{record['run']}`；训练曲线：`{record['run']}/results.csv`；模型：`{record['run']}/weights/best.pt`。这些路径仅服务器保存，不要求云端存在。", ""]
        if record["error"]:
            lines += [f"停止原因：{safe(record['error'])}", ""]
        for sweep in record["sweeps"]:
            lines += [f"评估：`{sweep['run']}`；范围：`{sweep['scope']}`；状态：`{sweep['status']}`。",
                      f"checkpoint SHA：`{sweep['protocol'].get('checkpoint_sha256', 'unknown')}`；GT SHA：`{sweep['protocol'].get('annotation_sha256', 'unknown')}`。", ""]
            for row in sweep["rows"]:
                if row["evaluation_status"] != "completed":
                    lines.append(f"- Q{row['queries']}：{row['evaluation_status']}；{safe(row['reason'])}")
                elif row.get("benchmark_status") not in (None, "completed"):
                    lines.append(f"- Q{row['queries']}计时：{row['benchmark_status']}；{safe(row['benchmark_reason'])}；FPS显示为—。")
    if data["collection_errors"]:
        lines += ["", "## 无法读取的实验", ""]
        lines += [f"- `{item['path']}`：{safe(item['error'])}" for item in data["collection_errors"]]
    lines += ["", "## 图表", "", "![训练内部mAP](plots/learning_curves.png)", "", "![Query精度](plots/query_accuracy.png)", "", "![Query计时](plots/query_latency.png)", ""]
    return "\n".join(lines)


def plots(data, output):
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "zhj-uav-5090-mpl"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    target = output / "plots"
    target.mkdir()
    main = rows(data, main=True)
    labels = sorted({row["run_id"] for row in main})
    for name in ("learning_curves", "query_accuracy", "query_latency"):
        fig, axes = plt.subplots(1, 2 if name == "query_accuracy" else 1, figsize=(11, 4), squeeze=False)
        ax = axes[0, 0]
        if name == "learning_curves":
            for record in data["records"]:
                if record["role"] == "paired_150e_screen" and record["curve"]:
                    ax.plot([r["epoch"] for r in record["curve"]], [100 * r[BEST_KEY] if r.get(BEST_KEY) is not None else math.nan for r in record["curve"]], label=record["run_id"])
            ax.set(xlabel="Epoch", ylabel="Internal validation mAP50:95 (%)")
        else:
            for label in labels:
                selected = sorted((r for r in main if r["run_id"] == label), key=lambda r: r["queries"])
                for index, key in enumerate(("AP", "APs") if name == "query_accuracy" else ("median_ms",)):
                    current = axes[0, index]
                    current.plot([r["queries"] for r in selected],
                                 [r[key] * (100 if key in METRICS else 1) if r[key] is not None else math.nan for r in selected], marker="o", label=label)
                    current.set(xlabel="Native decoder queries", ylabel=key + (" (%)" if key in METRICS else " (ms)"))
        for current in axes.flat:
            current.grid(alpha=.25)
            if current.get_legend_handles_labels()[0]:
                current.legend(fontsize=7)
            else:
                current.text(.5, .5, "No completed results", ha="center", transform=current.transAxes)
        fig.tight_layout()
        fig.savefig(target / f"{name}.png", dpi=150)
        plt.close(fig)


def excel(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    sheet = book.active
    sheet.title = "Query results"
    sheet.append(list(FIELDS))
    for row in rows(data):
        sheet.append([row.get(key) for key in FIELDS])
    for index, key in enumerate(FIELDS, 1):
        sheet.cell(1, index).font = Font(bold=True)
        if key in METRICS or key in ("coverage_05", "class_correct_coverage_05"):
            for column in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in column:
                    cell.number_format = "0.00%"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    training = book.create_sheet("Training status")
    keys = ("experiment", "training_status", "role", "epochs", "requested_epochs", "best_epoch", "best_internal_map", "error")
    training.append(list(keys))
    for record in data["records"]:
        training.append([record.get(key) for key in keys])
    stages = book.create_sheet("Stage latency")
    stages.append(list(stage_latency.FIELDS))
    for row in stage_latency.rows(data):
        stages.append([row[key] for key in stage_latency.FIELDS])
        stages.cell(stages.max_row, stage_latency.FIELDS.index("fraction_of_staged_forward") + 1).number_format = "0.00%"
    stages.freeze_panes = "A2"
    stages.auto_filter.ref = stages.dimensions
    book.save(path)


def generate_reports(runs_dir, output, *, data=None):
    """Serialize refreshes and atomically replace every generated artifact."""
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    with (output / ".report.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        data = collect(runs_dir) if data is None else data
        if data.get("schema") != "5090.experiments-report.v1":
            raise ValueError("Unsupported report snapshot schema")
        with tempfile.TemporaryDirectory(prefix=".5090-report-", dir=output) as directory:
            stage = Path(directory)
            (stage / "experiments.json").write_text(json.dumps(data, indent=2, sort_keys=True, allow_nan=False) + "\n")
            with (stage / "experiments.csv").open("w", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(rows(data))
            with (stage / "stage_latency.csv").open("w", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=stage_latency.FIELDS)
                writer.writeheader()
                writer.writerows(stage_latency.rows(data))
            plots(data, stage)
            excel(data, stage / "experiments.xlsx")
            (stage / "experiments.md").write_text(markdown(data))
            artifacts = sorted(path for path in stage.rglob("*") if path.is_file())
            (stage / "MANIFEST.sha256").write_text("".join(sha(path) + "  " + path.relative_to(stage).as_posix() + "\n"
                                                         for path in artifacts))
            for path in [*artifacts, stage / "MANIFEST.sha256"]:
                if path.is_file():
                    target = output / path.relative_to(stage)
                    target.parent.mkdir(parents=True, exist_ok=True)
                    os.replace(path, target)
        return data


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runs-dir", type=Path, default=ROOT / "runs/decoder_budget_5090")
    parser.add_argument("--output", type=Path, default=ROOT / "reports/5090")
    parser.add_argument("--from-json", type=Path, help="Render an archived snapshot without server-only assets")
    args = parser.parse_args()
    data = generate_reports(args.runs_dir, args.output, data=read(args.from_json) if args.from_json else None)
    print(json.dumps({"output": str(args.output.resolve()), "training_runs": len(data["records"]),
                      "full_validation_rows": len(rows(data, main=True)), "collection_errors": data["collection_errors"]}))


if __name__ == "__main__":
    main()
